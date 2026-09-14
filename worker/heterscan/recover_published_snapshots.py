"""Recover legacy normalized history from its preserved, published workbook."""
from __future__ import annotations

import json
from datetime import date, datetime
from io import BytesIO
from urllib.parse import quote

from openpyxl import load_workbook

from .supabase import SupabaseRepository

FIELDS = {
    "address": "כתובת", "application_number": "מספר בקשה", "building_file_number": "מספר תיק בניין",
    "block_number": "גוש", "parcel_number": "חלקה", "application_type": "סוג בקשה",
    "work_description": "תיאור עבודה", "submission_date": "תאריך הגשה", "approval_date": "תאריך אישור",
    "permit_number": "מספר היתר", "permit_issue_date": "תאריך הפקת היתר",
    "permit_status_original": "סטטוס מקורי במקור", "permit_confidence": "רמת אמינות", "source_url": "קישור מקור",
}
MISSING = {"", "לא ידוע", "טרם הופק", "—"}


def _value(value):
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return None if value is None or str(value).strip() in MISSING else str(value)


def _rows(sheet):
    values = iter(sheet.iter_rows(values_only=True))
    headers = next(values)
    return [dict(zip(headers, row)) for row in values if any(value is not None for value in row)]


def prepare_recovery(payload: bytes, links: list[dict], report_path: str):
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        rows = _rows(workbook["בקשות והיתרים"])
        by_number = {}
        for row in rows:
            number = _value(row["מספר בקשה"])
            if number is None or number in by_number:
                raise ValueError("Published application numbers are missing or ambiguous")
            by_number[number] = row
        issued = {_value(row["מספר בקשה"]) for row in _rows(workbook["היתרים שנמצאו"])}
        approved = {_value(row["מספר בקשה"]) for row in _rows(workbook["בקשות שאושרו"])}
        summary = _rows(workbook["סיכום"])[0]
        if len(rows) != int(summary["בקשות שנמצאו"]) or len(issued) != int(summary["היתרים שנמצאו"]):
            raise ValueError("Published workbook counts do not reconcile")
        recovery = []
        for link in links:
            if link["result_snapshot_source"] != "legacy_current":
                continue
            snapshot = dict(link["result_snapshot"])
            source = by_number.get(str(snapshot.get("application_number")))
            if source is None:
                raise ValueError("A legacy application is absent from the published workbook")
            snapshot.update({key: _value(source[label]) for key, label in FIELDS.items()})
            number = snapshot["application_number"]
            snapshot["is_permit_issued"] = number in issued
            snapshot["is_approved"] = number in approved
            # These normalized values are report-derived, not reconstructed raw evidence.
            snapshot.pop("raw_data", None)
            snapshot.pop("content_hash", None)
            snapshot.pop("evidence", None)
            snapshot["legacy_report_path"] = report_path
            recovery.append({"application_id": link["application_id"], "snapshot": snapshot})
        return recovery, len(rows), len(issued)
    finally:
        workbook.close()


def main() -> None:
    repository = SupabaseRepository()
    failures = []
    try:
        runs = repository._get_all(
            "run_overview?status=not.in.(created,dispatching,running,safely_stopped)&report_path=not.is.null"
            "&select=id,report_path&order=created_at"
        )
        for run in runs:
            run_id, path = run["id"], run["report_path"]
            try:
                if not path.startswith(f"{run_id}/"):
                    raise ValueError("Unexpected report path")
                links = repository._get_all(
                    f"run_applications?run_id=eq.{quote(run_id)}"
                    "&select=application_id,result_snapshot,result_snapshot_source"
                )
                if not any(row["result_snapshot_source"] == "legacy_current" for row in links):
                    continue
                response = repository.client.get(
                    f"{repository.url}/storage/v1/object/authenticated/reports/{quote(path, safe='/')}"
                )
                response.raise_for_status()
                rows, applications, permits = prepare_recovery(response.content, links, path)
                result = repository._rest("POST", "rpc/recover_published_run_snapshots", json={
                    "p_run_id": run_id, "p_rows": rows, "p_report_path": path,
                    "p_applications": applications, "p_permits": permits,
                }).json()
                print(json.dumps({"run_id": run_id, "recovered": result, "applications": applications, "permits": permits}))
            except Exception as error:
                failures.append(run_id)
                print(json.dumps({"run_id": run_id, "error": str(error)}))
    finally:
        repository.close()
    if failures:
        raise SystemExit(f"Recovery requires review for {len(failures)} runs; other runs preserved")


if __name__ == "__main__":
    main()
