"""Synthetic offline fixtures: no source, challenge service or session access."""

from datetime import date
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
import threading
from urllib.parse import parse_qs

import httpx
import pytest
from openpyxl import load_workbook

from heterscan.adapters.complot import ComplotAdapter
from heterscan.domain import AdapterRateLimited, AdapterReviewRequired, SearchUnit
from heterscan.http import AdaptiveRateLimiter, PublicHttpClient, _challenge_diagnostics
from heterscan.reporting import build_report
from heterscan.runner import _collect_wave, _final_status

FIXTURES = Path(__file__).parent / "fixtures" / "challenges"
DATES = (date(2026, 1, 1), date(2026, 1, 31))
UNIT = SearchUnit("fixture-unit", "fixture-run", 1, "request:20269999",
                  {"mode": "request", "requestNumber": "20269999"})


def fixture(name):
    return (FIXTURES / name).read_text(encoding="utf-8")


@pytest.fixture
def client():
    value = PublicHttpClient(delay_seconds=0)
    value.client.close()
    yield value
    value.close()


@pytest.mark.parametrize("body,status,classification", [
    (fixture("widget.html"), 200, "visible_challenge"),
    ("<p>CAPTCHA required</p>", 200, "visible_challenge"),
    ("<p>אני לא רובוט</p>", 200, "visible_challenge"),
    ('<div class="g-recaptcha"></div>', 200, "visible_challenge"),
    ('<iframe src="https://www.google.com/recaptcha/api2/anchor?size=normal"></iframe>',
     200, "visible_challenge"),
    (fixture("f5.html"), 200, "f5_rejection"),
    (fixture("f5.html"), 404, "f5_rejection"),
    ("<p>Access denied</p>", 200, "source_blocked"),
    ("<p>The requested URL was rejected. Your support ID is FIXTURE.</p>", 200, "source_blocked"),
    (fixture("detail.html"), 403, "http_blocked"),
    (fixture("detail.html"), 401, "http_blocked"),
    ("x" * 5100 + fixture("widget.html"), 200, "visible_challenge"),
])
def test_blocked_responses_stop_without_retry_and_latch_origin(client, body, status, classification):
    calls = []

    def handler(request):
        calls.append(request)
        return httpx.Response(status, text=body, headers={"Set-Cookie": "DO_NOT_LOG_COOKIE"}, request=request)

    client.client = httpx.Client(transport=httpx.MockTransport(handler))
    with pytest.raises(AdapterReviewRequired) as raised:
        client.request("GET", "https://example.test/detail")
    diagnostics = raised.value.diagnostics
    assert diagnostics["classification"] == classification
    assert diagnostics["http_status"] == status
    assert diagnostics["source_origin"] == "example.test"
    assert diagnostics["observed_at"]
    assert "DO_NOT_LOG" not in str(diagnostics)
    with pytest.raises(AdapterReviewRequired) as skipped:
        client.request("GET", "https://example.test/another-detail")
    assert skipped.value.diagnostics["request_skipped"] is True
    assert skipped.value.diagnostics["http_status"] is None
    assert skipped.value.diagnostics["visible_challenge"] is None
    assert skipped.value.diagnostics["blocking_evidence"]["classification"] == classification
    assert len(calls) == 1


@pytest.mark.parametrize("hidden", [
    '<div hidden><div class="g-recaptcha">CAPTCHA</div></div>',
    '<div style="display: none"><div class="g-recaptcha">CAPTCHA</div></div>',
    '<div aria-hidden="true">CAPTCHA</div>',
    '<div class="g-recaptcha" data-size="invisible"></div>',
    '<template><div class="g-recaptcha">CAPTCHA</div></template>',
    '<!-- CAPTCHA -->',
    '<textarea>captcha token</textarea>',
    '<iframe src="https://[malformed"></iframe>',
    '<div class="grecaptcha-badge"><iframe src="https://www.google.com/recaptcha/api2/anchor?size=invisible"></iframe>protected by reCAPTCHA</div>',
])
def test_hidden_markup_is_not_a_confirmed_visible_challenge(hidden):
    evidence = _challenge_diagnostics("<html><body>" + hidden + "</body></html>")
    assert evidence["visible_challenge"] is False


def configure_adapter(detail, status=200):
    adapter = ComplotAdapter("8400", "רחובות", {"site_id": "22"})
    adapter.client.rate_limiter = None  # Fixtures must not wait or call the coordinator.
    adapter.client.client.close()
    calls = []

    def handler(request):
        calls.append(request)
        is_summary = parse_qs(request.url.query.decode()).get("prgname") == ["GetBakashotByNumber"]
        return httpx.Response(200 if is_summary else status,
                              text=fixture("summary.html") if is_summary else detail,
                              request=request)

    adapter.client.client = httpx.Client(transport=httpx.MockTransport(handler))
    # No replacement of adapter.collect or validator: exercise their real flow.
    return adapter, calls


@pytest.mark.parametrize("with_script", [False, True])
def test_valid_detail_is_parsed_and_validator_decides(with_script):
    body = fixture("detail.html")
    if with_script:
        body = body.replace("<body>", "<body>" + fixture("script-only.html"))
    adapter, calls = configure_adapter(body)
    try:
        [(_, records, error)] = _collect_wave(adapter, [UNIT], *DATES)
        assert error is None and len(records) == 1
        record = records[0]
        assert record.details_available is True
        assert record.application_number == "20269999"
        assert record.submission_date == date(2026, 1, 14)
        assert record.raw_data["validator"]["status"] == "passed_consistency_checks"
        evidence = record.raw_data["source_diagnostics"]
        assert evidence["parser_result"] == "parsed"
        assert evidence["visible_challenge"] is False
        assert evidence["challenge_script_marker"] is with_script
        assert "DO_NOT_LOG" not in str(evidence)
        assert len(calls) == 2
    finally:
        adapter.close()


def test_valid_issued_detail_with_script_and_branding_preserves_verified_permit():
    body = fixture("detail.html").replace("</table>",
        '<tr><td>מספר היתר</td><td>FIXTURE-PERMIT</td></tr>'
        '<tr><td>תאריך היתר</td><td>20/01/2026</td></tr></table>')
    body = body.replace("<body>", "<body>" + fixture("script-only.html")
        + '<div class="grecaptcha-badge">protected by reCAPTCHA</div>')
    adapter, _ = configure_adapter(body)
    try:
        [(_, records, error)] = _collect_wave(adapter, [UNIT], *DATES)
        assert error is None
        row = records[0].to_database(run_id="fixture-run", identity_key="fixture", content_hash="fixture")
        assert row["permit_verification"] == "verified_issued"
        assert row["permit_number"] == "FIXTURE-PERMIT"
        assert row["permit_issue_date"] == "2026-01-20"
        diagnostics = records[0].raw_data["source_diagnostics"]
        assert diagnostics["challenge_script_marker"] is True
        assert diagnostics["challenge_badge"] is True
        assert diagnostics["visible_challenge"] is False
    finally:
        adapter.close()


@pytest.mark.parametrize("body,status,parser_result", [
    (fixture("widget.html"), 200, None),
    (fixture("detail.html"), 403, None),
    (fixture("script-only.html"), 200, "detail_missing"),
    ("", 200, "detail_missing"),
    ("   ", 200, "detail_missing"),
    (fixture("detail.html").replace('<div id="info-main"><table>', '<div><table>'), 200, "detail_missing"),
    ('<div id="result-title-div-id">מספר הבקשה: 20269999 תאריך הגשה: 14/01/2026</div>'
     '<div id="info-main"></div>' + fixture("script-only.html"), 200, "detail_missing"),
    ('<div id="info-main"></div><div id="result-title-div-id">changed layout</div>'
     + fixture("script-only.html"), 200, "parser_mismatch"),
    (fixture("detail.html").replace("20269999", "20269998"), 200, "identity_mismatch"),
])
def test_unreadable_detail_preserves_summary_unknown_permit_and_review(body, status, parser_result):
    adapter, calls = configure_adapter(body, status)
    try:
        [(_, records, error)] = _collect_wave(adapter, [UNIT], *DATES)
        assert isinstance(error, AdapterReviewRequired)
        assert len(records) == 1
        record = records[0]
        assert record.details_available is False
        assert record.application_number == "20269999"
        assert record.address == "רחוב הבדיקה 11"
        assert record.submission_date == date(2026, 1, 14)
        assert record.permit_number is None and not record.is_permit_issued
        evidence = record.raw_data["source_diagnostics"]
        assert evidence["http_status"] == status
        if parser_result:
            assert evidence["parser_result"] == parser_result
            assert evidence["visible_challenge"] is False
        else:
            assert evidence["classification"] in {"visible_challenge", "http_blocked"}
        assert len(calls) == 2
        assert _final_status({"units_requires_review": 1, "units_failed": 0, "applications_found": 1}) == "requires_review"
        row = record.to_database(run_id="fixture-run", identity_key="fixture", content_hash="fixture")
        assert row["permit_verification"] == "unknown"
        payload, _ = build_report({"city_name": "רחובות", "date_from": "2026-01-01",
                                   "date_to": "2026-01-31", "status": "requires_review"}, [row], [])
        workbook = load_workbook(BytesIO(payload))
        assert workbook["סיכום"]["E2"].value == 1
        assert workbook["סיכום"]["F2"].value == 0
        assert "היתר שלא אומת אינו היתר שלא הופק" in workbook["סיכום"]["H2"].value
        assert workbook["היתרים שנמצאו"].max_row == 1
        assert workbook["בקשות והיתרים"]["J2"].value == "לא ידוע"
        assert workbook["בקשות והיתרים"]["L2"].value == "פרטים חלקיים — נדרשת בדיקה"
    finally:
        adapter.close()


@pytest.mark.parametrize("body", ["", fixture("script-only.html"), "<html><body>changed layout</body></html>"])
def test_legacy_search_does_not_convert_empty_or_script_only_to_zero(body):
    adapter = ComplotAdapter("8400", "רחובות", {"site_id": "22"})
    try:
        with pytest.raises(AdapterReviewRequired) as raised:
            adapter._list_rows(body)
        assert raised.value.diagnostics["parser_result"] == "parser_mismatch"
    finally:
        adapter.close()


def test_retry_after_prevents_second_transport_call(client):
    calls = []
    limiter = AdaptiveRateLimiter(requests_per_second=100)
    client.rate_limiter = limiter

    def handler(request):
        calls.append(request)
        return httpx.Response(429, headers={"Retry-After": "120"}, request=request)

    client.client = httpx.Client(transport=httpx.MockTransport(handler))
    with pytest.raises(AdapterRateLimited) as raised:
        client.request("GET", "https://example.test/detail")
    assert raised.value.retry_after_seconds == 120
    with pytest.raises(AdapterRateLimited):
        client.request("GET", "https://example.test/detail")
    assert len(calls) == 1


def test_call_already_waiting_for_slot_is_skipped_after_challenge(client, monkeypatch):
    limiter = AdaptiveRateLimiter(requests_per_second=1000)
    client.rate_limiter = limiter
    first_in_transport = threading.Event()
    second_waiting = threading.Event()
    lock = threading.Lock()
    entries = 0
    transport_calls = []
    original_wait = limiter.wait

    def wait():
        nonlocal entries
        with lock:
            entries += 1
            if entries == 2:
                second_waiting.set()
        original_wait()

    def handler(request):
        transport_calls.append(request)
        first_in_transport.set()
        assert second_waiting.wait(timeout=2)
        return httpx.Response(200, text=fixture("widget.html"), request=request)

    monkeypatch.setattr(limiter, "wait", wait)
    client.client = httpx.Client(transport=httpx.MockTransport(handler))
    with ThreadPoolExecutor(max_workers=2) as pool:
        first = pool.submit(client.request, "GET", "https://example.test/first")
        assert first_in_transport.wait(timeout=2)
        second = pool.submit(client.request, "GET", "https://example.test/second")
        with pytest.raises(AdapterReviewRequired):
            first.result(timeout=3)
        with pytest.raises(AdapterReviewRequired) as skipped:
            second.result(timeout=3)
    assert skipped.value.diagnostics["request_skipped"] is True
    assert len(transport_calls) == 1
    assert limiter._slots.acquire(blocking=False)  # Neither error leaked its lane.
    limiter._slots.release()
