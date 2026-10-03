"""Public-summary mode is selected before collection, never after a denial."""
from datetime import date
from io import BytesIO
from pathlib import Path
from urllib.parse import parse_qs

import httpx
import pytest
from openpyxl import load_workbook

from heterscan.adapters.complot import ComplotAdapter
from heterscan.domain import AdapterRateLimited, AdapterReviewRequired, SearchUnit
from heterscan.reporting import build_report
from heterscan.runner import _collect_wave, _final_status

FIXTURE = Path(__file__).parent / "fixtures/challenges/summary.html"
UNIT = SearchUnit("summary-unit", "summary-run", 1, "request:20269999",
                  {"mode": "request", "requestNumber": "20269999"})
JANUARY = (date(2026, 1, 1), date(2026, 1, 31))


def configured_adapter(status=200, body=None):
    adapter = ComplotAdapter("8400", "רחובות", {"site_id": "22", "collection_mode": "public_summary"})
    adapter.client.client.close()
    calls = []

    def handler(request):
        calls.append(request)
        # A guarded detail route must not even be requested in this mode.
        assert parse_qs(request.url.query.decode()).get("prgname") == ["GetBakashotByNumber"]
        return httpx.Response(status, text=body if body is not None else FIXTURE.read_text(encoding="utf-8"),
                              headers={"Retry-After": "120"} if status == 429 else {}, request=request)

    adapter.client.client = httpx.Client(transport=httpx.MockTransport(handler))
    return adapter, calls


def test_summary_only_uses_one_public_request_and_preserves_field_provenance_in_report():
    adapter, calls = configured_adapter()
    try:
        [(_, records, error)] = _collect_wave(adapter, [UNIT], *JANUARY)
        assert isinstance(error, AdapterReviewRequired)
        assert len(records) == 1 and len(calls) == 1
        record = records[0]
        assert record.details_available is False
        assert record.application_number == "20269999"
        assert record.raw_data["source_diagnostics"]["detail_requested"] is False
        assert record.raw_data["source_diagnostics"]["detail_state"] == "not_requested"
        assert record.raw_data["validator"]["status"] == "requires_review"
        assert set(record.raw_data["field_provenance"]) == {
            "application_number", "address", "submission_date", "building_file_number", "block_number", "parcel_number",
        }
        row = record.to_database(run_id="summary-run", identity_key="fixture", content_hash="fixture")
        assert row["permit_verification"] == "unknown"
        assert _final_status({"units_requires_review": 1, "units_failed": 0, "applications_found": 1}) == "requires_review"
        run = {"city_name": "רחובות", "date_from": "2026-01-01", "date_to": "2026-01-31",
               "status": "requires_review", "configuration_snapshot": {"collectionMode": "public_summary"}}
        payload, _ = build_report(run, [row], [{"status": "requires_review"}])
        workbook = load_workbook(BytesIO(payload))
        assert workbook["סיכום"]["F2"].value == 0
        assert workbook["סיכום"]["K2"].value == "סיכומי בקשות בלבד"
        assert workbook["סיכום"]["L2"].value == 1
        assert workbook["היתרים שנמצאו"].max_row == 1
        assert workbook["בקשות והיתרים"]["J2"].value == "לא ידוע"
        assert workbook["בקשות והיתרים"]["L2"].value == "סיכום בקשה בלבד — מצב היתר לא ידוע"
        sheet = workbook["מקור שדות הסיכום"]
        assert sheet.max_row == 7
        assert all(sheet.cell(i, 6).hyperlink.target == record.source_url for i in range(2, 8))
    finally:
        adapter.close()


@pytest.mark.parametrize("status,body,exception", [
    (403, "blocked", AdapterReviewRequired),
    (200, '<div class="g-recaptcha"></div>', AdapterReviewRequired),
    (200, "", AdapterReviewRequired),
    (429, "rate limited", AdapterRateLimited),
])
def test_summary_mode_still_stops_when_public_search_is_unreadable(status, body, exception):
    adapter, calls = configured_adapter(status, body)
    try:
        with pytest.raises(exception):
            adapter.collect(UNIT, *JANUARY)
        assert len(calls) == 1
    finally:
        adapter.close()


def test_out_of_range_summary_is_filtered_without_any_detail_request():
    adapter, calls = configured_adapter()
    try:
        assert adapter.collect(UNIT, date(2026, 2, 1), date(2026, 2, 28)) == []
        assert len(calls) == 1
    finally:
        adapter.close()


def test_invalid_mode_and_legacy_strategy_fail_before_transport():
    with pytest.raises(ValueError):
        ComplotAdapter("8400", "רחובות", {"site_id": "22", "collection_mode": "captcha_solver"})
    adapter, calls = configured_adapter()
    try:
        with pytest.raises(AdapterReviewRequired):
            adapter.collect(SearchUnit("u", "r", 1, "street:1", {"streetCode": "1"}), *JANUARY)
        assert calls == []
    finally:
        adapter.close()
